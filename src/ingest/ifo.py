"""ifo Business Climate Germany (ifo Institute, public Excel "gsk-e-YYYYMM.xlsx", no key).

Sheet 1: Germany (industry and trade) climate / situation / expectations as index (2015=100, SA) and
balances (SA) since 2005-01, ifo Business Uncertainty and the Business Cycle Traffic Lights probability.
Sheet 'Sectors': the same by sector (manufacturing, services, trade, wholesaling, retailing, construction)
since 1991-01. Western-Germany history 1960-1990 is a separate Excel (ifo_Business_Survey_Historical_TimeSeries.xlsx);
older series only on request by e-mail (umfragedaten@ifo.de) -> not collected.

The file name carries the release month; the script tries the current month and goes back up to 6 months.
Run: .venv/bin/python -m src.ingest.ifo
Raw: $FIN_DATA_ROOT/raw/ifo/gsk-e-YYYYMM.xlsx   Curated: curated/ifo/ifo_*.parquet + _catalog.csv
"""
from __future__ import annotations

import io
import re

import openpyxl
import pandas as pd

from src.ingest._europe_common import CatalogWriter, add_series, cached_get, log, raw_dir

SOURCE = "ifo"
URL = "https://www.ifo.de/sites/default/files/secure/timeseries/gsk-e-{ym}.xlsx"
PAGE = "https://www.ifo.de/en/umfragen/time-series"
LICENSE = ("ifo Institute: published for free download on ifo.de; cite 'ifo Business Climate Germany, ifo Institute'. "
           "Redistribution terms not stated (treat as research use, non-commercial)")
SECT = {"Industry and Trade (Manufacturing, Trade, and Construction)": "it", "Manufacturing": "mfg", "Services": "serv",
        "Trade": "trade", "Wholesaling": "whole", "Retailing": "retail", "Construction": "constr"}
SECT_EN = {"it": "industry and trade", "mfg": "manufacturing", "serv": "services", "trade": "trade (wholesale+retail)",
           "whole": "wholesaling", "retail": "retailing", "constr": "construction"}


def get_file() -> tuple[bytes, str]:
    today = pd.Timestamp.today()
    for k in range(0, 7):
        ym = (today - pd.DateOffset(months=k)).strftime("%Y%m")
        try:
            c = cached_get(URL.format(ym=ym), raw_dir(SOURCE) / f"gsk-e-{ym}.xlsx", min_bytes=20000, retries=1, allow_404=True, max_age_hours=24 * 7)
        except Exception as e:  # noqa: BLE001
            log.info("ifo %s not available: %s", ym, e)
            continue
        if c and c[:2] == b"PK":
            return c, ym
    raise RuntimeError("ifo: no gsk-e file found in the last 7 months")


def rows(ws):
    out = {}
    for r in ws.iter_rows(min_row=9, values_only=True):
        m = re.fullmatch(r"\s*(\d{2})/(\d{4})\s*", str(r[0] or ""))
        if m:
            out[pd.Timestamp(int(m.group(2)), int(m.group(1)), 1)] = r
    return pd.DataFrame.from_dict(out, orient="index").sort_index()


def main() -> pd.DataFrame:
    content, ym = get_file()
    wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
    cw = CatalogWriter(SOURCE, url=PAGE, license=LICENSE)
    note = f"ifo file gsk-e-{ym}.xlsx. Monthly, dated at month start (survey month)."
    df = rows(wb.worksheets[0])
    spec = [(1, "climate_idx", "Business Climate (index 2015=100, SA)", "index 2015=100"),
            (2, "situation_idx", "Current business situation (index 2015=100, SA)", "index 2015=100"),
            (3, "expectations_idx", "Business expectations (index 2015=100, SA)", "index 2015=100"),
            (4, "climate_bal", "Business Climate (balance, SA)", "balance (pp)"),
            (5, "situation_bal", "Current business situation (balance, SA)", "balance (pp)"),
            (6, "expectations_bal", "Business expectations (balance, SA)", "balance (pp)"),
            (7, "uncertainty", "ifo Business Uncertainty (original values after 04/2019)", "index"),
            (8, "traffic_light_prob", "ifo Business Cycle Traffic Lights: probability of economic expansion", "%")]
    for c, tag, nm, units in spec:
        add_series(cw, f"ifo_de_{tag}_m", pd.to_numeric(df[c], errors="coerce"), name=f"ifo Germany: {nm}", category="growth",
                   frequency="M", units=units, country="DE", transform_notes=note + " Sheet 'ifo Business Climate' (since 2005-01).")
    ws = wb["Sectors"]
    groups = next(ws.iter_rows(min_row=6, max_row=6, values_only=True))
    names = next(ws.iter_rows(min_row=8, max_row=8, values_only=True))
    ds = rows(ws)
    cur = None
    for c in range(1, len(groups)):
        if groups[c]:
            cur = SECT.get(str(groups[c]).strip())
        if cur is None:
            continue
        nm = str(names[c]).strip().lower().replace("business ", "")
        is_idx = (cur == "it" and c <= 3)
        tag = f"{cur}_{nm}_{'idx' if is_idx else 'bal'}"
        add_series(cw, f"ifo_de_sect_{tag}_m", pd.to_numeric(ds[c], errors="coerce"),
                   name=f"ifo Germany, {SECT_EN[cur]}: business {nm} ({'index 2015=100' if is_idx else 'balance'}, SA)",
                   category="growth", frequency="M", units="index 2015=100" if is_idx else "balance (pp)", country="DE",
                   transform_notes=note + " Sheet 'Sectors' (since 1991-01; services/trade shorter).")
    return cw.close()


if __name__ == "__main__":
    main()
